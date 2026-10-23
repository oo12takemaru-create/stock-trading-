# -*- coding: utf-8 -*-
"""投資部門別売買状況の週次履歴を再構築する → docs/investor_hist.json

■ なぜ作るか
  investor_flow.json は「今週と前週」の2点しか持っていない。
  AI検索から来る質問は「誰が買っているか」に集中しているのに、2点では答えられない。
  JPXの年別アーカイブを辿れば週次の履歴を作れる。

■ 出所
  JPX「投資部門別売買状況」の年別アーカイブ
  https://www.jpx.co.jp/markets/statistics-equities/investor-type/00-00-archives-NN.html
  ファイル名 stock_val_1_YYMMWW.xls が「YY年MM月第WW週」を表す。
  **アーカイブは2016年までしか遡れない**（archives-11 以降は404）。
  取れない年は取らない。埋めない。

■ 市場区分が2022年4月に変わっている
  2022-04-04 の再編で「東証1部」→「プライム」。Excelのシート名も TSE 1st → TSE Prime。
  構成銘柄が違うので、**黙って1本の系列として繋がない**。
  JSONの meta に変更日を持たせ、画面でも分かるようにする。

■ 2026年9月第3週（9/14〜）からファイル形式が変わった
  旧: stock_val_1_YYMMWW.xls（シート "TSE Prime"、主体が行）
  新: stock_1_w_YYYYMMDD_YYYYMMDD.xlsx（1シートだけ・横持ち。主体が列、市場が行）
  ファイル名に週の開始日・終了日がそのまま入っている。
  新形式の読み方は parse_new() にまとめ、investor_flow.py / investor_hist_append.py も
  ここを使う（3本で読み方がずれないように）。
  ★週番号の数え方も変わった★ 新形式のA列コード（例 2026095）は「月曜が属する月」で数える。
  9/28〜10/2 は旧形式なら「10月第1週」だったが、新形式は「9月第5週」。JPX自身の表記に従う。

■ 旧形式は「前週｜今週」の2ブロックが横に並んでいる（2026-10-23 に判明・修正）
  11行目あたりに「08/31～09/04 | 09/07～09/11」の週ラベルがあり、各行は
  [前週 金額・比率・差引] [今週 金額・比率・差引] の順。以前は行の「最初の大きな数」を
  取っていたため**前週の値を拾っていた**（旧形式の履歴が丸ごと1週ずれていた）。
  いまは週ラベル行から今週ブロックの列を特定して読む（parse_old_df）。
  ファイルの週（ラベル・ファイル名）と選んだブロックの日付が一致しなければ止める。
  investor_flow.py / precompute/score3_hist.py も parse_old_df を使う（読み方を1本化）。

■ パースの作法（カンマ入り文字列・千円→億円）は investor_flow.py と合わせる
"""
import io
import json
import re
import sys
import time
import urllib.request
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

JST = timezone(timedelta(hours=9))
HERE = Path(__file__).parent
CACHE = HERE / "investor_cache"
OUT = HERE / "docs" / "investor_hist.json"
BASE = "https://www.jpx.co.jp"
INDEX = "/markets/statistics-equities/investor-type/index.html"
ARCHIVE = "/markets/statistics-equities/investor-type/00-00-archives-{n:02d}.html"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}

# 市場区分の変更日（東証1部 → プライム）
REORG = "2022-04-04"

CATEGORIES = [
    ("海外投資家", "foreigners",  "海外投資家"),
    ("個",         "individuals", "個人投資家"),
    ("信託銀行",   "trust_banks", "信託銀行（年金など）"),
    ("投資信託",   "inv_trusts",  "投資信託"),
    ("事業法人",   "business",    "事業法人（自社株買いなど）"),
    ("自己",       "proprietary", "証券自己（プロップ）"),
]


def http(url, tries=3):
    for i in range(tries):
        try:
            return urllib.request.urlopen(
                urllib.request.Request(url, headers=UA), timeout=60).read()
        except Exception as e:
            if i == tries - 1:
                raise
            print(f"  再試行{i+1} {url}: {e}", file=sys.stderr)
            time.sleep(3 * (i + 1))


def links():
    """最新ページ＋年別アーカイブから、週次Excelの URL を全部集める
    キーは 旧形式=YYMMWW（6桁）、新形式=w+開始日YYYYMMDD（並び順は週の開始日で後から直す）"""
    found = {}
    pages = [INDEX] + [ARCHIVE.format(n=n) for n in range(0, 11)]
    for p in pages:
        try:
            s = http(BASE + p).decode("utf-8", "replace")
        except Exception as e:
            print(f"  ページ取得失敗 {p}: {e}", file=sys.stderr)
            continue
        n = 0
        for m in re.finditer(r'href="([^"]*stock_val_1_(\d{6})\.xls)"', s):
            found.setdefault(m.group(2), BASE + m.group(1))
            n += 1
        for m in re.finditer(r'href="([^"]*' + NEW_FILE_RE.pattern + ')"', s):
            found.setdefault("w" + m.group(2), BASE + m.group(1))
            n += 1
        print(f"  {p.split('/')[-1]:28s} {n}件")
        time.sleep(0.4)
    return found


def sheet_for(xl):
    """東証1部（〜2022年3月）と プライム（2022年4月〜）でシート名が違う"""
    for want in ("TSE Prime", "TSE 1st", "TSE 1st section"):
        for s in xl.sheet_names:
            if s.strip().lower() == want.lower():
                return s
    return xl.sheet_names[0]


def market_of(sheet):
    """シート名から市場区分を決める。日付ではなくシート名で判定する
    （JPXが区分を変えたらシート名も変わる。日付でハードコードすると次の再編で壊れる）"""
    s = (sheet or "").strip().lower()
    if s.startswith("tse prime"):
        return "prime"
    if s.startswith("tse 1st"):
        return "tse1"
    return "unknown"


def week_dates(label, code):
    """週ラベル末尾の "( 9/1 - 9/4 )" と ファイル名の YYMMWW から開始日・終了日を作る"""
    y = 2000 + int(code[:2])
    m = re.search(r"\(\s*(\d{1,2})/(\d{1,2})\s*-\s*(\d{1,2})/(\d{1,2})\s*\)", label or "")
    if not m:
        return None, None
    m1, d1, m2, d2 = (int(x) for x in m.groups())
    try:
        # 年をまたぐ週（12/29 - 1/4 のような形）
        return date(y, m1, d1).isoformat(), date(y + 1 if m2 < m1 else y, m2, d2).isoformat()
    except ValueError:
        return None, None


# ===== 新形式（2026年9月第3週〜 stock_1_w_YYYYMMDD_YYYYMMDD.xlsx） =====
NEW_FILE_RE = re.compile(r"stock_1_w_(\d{8})_(\d{8})\.xlsx")

# 旧形式の行の意味に合わせた、新形式での列の取り方。
# 見出しは4段（自己/委託 → 個人・海外投資家・法人… → 現金・信用・投資信託… → 金融機関の内訳）。
# 段ごとの日本語ラベル(p[0]〜p[3])で列グループを選び、該当グループの 売・買 を足す。
#   旧「自己計」      = 自己の 現金＋信用
#   旧「個人」        = 委託・個人の 現金＋信用
#   旧「海外投資家」  = 委託・海外投資家の 法人＋個人
#   旧「投資信託」「事業法人」「信託銀行」 = 委託・法人の内訳
# （旧形式 stock_val_1_260902.xls の内訳表 個人現金/個人信用・自己現金/自己信用・
#   海外投資家の法人/個人 を足すと、それぞれ 個人・自己計・海外投資家 に一致することを確認済み）
# n は足すグループの数（列の取り違えを検出するため数も確かめる）
NEW_SELECT = [
    ("foreigners",  2, lambda p: p[0] == "委託" and p[1] == "海外投資家"),
    ("individuals", 2, lambda p: p[0] == "委託" and p[1] == "個人"),
    ("trust_banks", 1, lambda p: p[0] == "委託" and p[1] == "法人" and p[3] == "信託銀行"),
    ("inv_trusts",  1, lambda p: p[0] == "委託" and p[1] == "法人" and p[2] == "投資信託"),
    ("business",    1, lambda p: p[0] == "委託" and p[1] == "法人" and p[2] == "事業法人"),
    ("proprietary", 2, lambda p: p[0] == "自己"),
]


def is_new_format(raw):
    """xlsx（zip）なら新形式。旧形式は xls（BIFF）"""
    return raw[:2] == b"PK"


def _ja(v):
    """「海外投資家 Foreigners」→「海外投資家」。見出しの日本語部分だけを取る"""
    if v is None:
        return ""
    return re.split(r"[\s\u3000\xa0/／]+", str(v).strip())[0]


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return None if v != v else float(v)          # NaN（pandas の空セル）は None
    t = str(v).replace(",", "").replace("△", "-").strip()
    return float(t) if re.fullmatch(r"-?\d+(\.\d+)?", t) else None


def parse_new(raw):
    """新形式の週次xlsxから 東証プライム・金額 の6主体を取る。単位は旧形式と同じ千円。

    列番号は決め打ちしない。結合セル（ws.merged_cells）を展開して各列の見出しの
    階層を作り、NEW_SELECT の条件で列グループを選ぶ。
    戻り値: {"code","market","sheet","flows":{key:{"sell","buy"}}(千円),"net":{key:億円}}
    """
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.worksheets[0]
    nrow, ncol = ws.max_row, ws.max_column

    grid = [[ws.cell(r, c).value for c in range(1, ncol + 1)] for r in range(1, nrow + 1)]
    for mr in ws.merged_cells.ranges:          # 結合セルは左上の値で埋める
        v = grid[mr.min_row - 1][mr.min_col - 1]
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                grid[r - 1][c - 1] = v

    # 「売 Sales / 買 Purchases / 差引 Balance / 合計 Total」の段を探す。その上4段が見出し
    sub_row = next((i for i, row in enumerate(grid)
                    if sum(1 for v in row if _ja(v) == "売") >= 6), None)
    if sub_row is None or sub_row < 4:
        raise ValueError("新形式: 売/買 の見出し段が見つからない")
    levels = grid[sub_row - 4:sub_row]

    # 列グループ: 見出し4段が同じ列のまとまり（売・買・差引・合計）
    groups = {}
    for c in range(ncol):
        sub = _ja(grid[sub_row][c])
        if sub not in ("売", "買", "差引", "合計"):
            continue
        path = tuple(_ja(levels[k][c]) for k in range(4))
        groups.setdefault(path, {})[sub] = c
    if not any(p[0] == "自己" for p in groups) or not any(p[0] == "委託" for p in groups):
        raise ValueError("新形式: 自己/委託 の見出しが見つからない")

    # 東証プライムの「金額 Value」行
    target, code, market = None, None, None
    for i in range(sub_row + 1, nrow):
        row = grid[i]
        mk = str(row[1] or "")
        if mk.startswith("東証プライム") and _ja(row[2]) == "金額":
            target, code, market = row, str(row[0] or "").strip(), re.sub(r"\s+", " ", mk)
            break
    if target is None:
        raise ValueError("新形式: 東証プライムの金額行が見つからない")

    flows, net = {}, {}
    for key, n_expect, cond in NEW_SELECT:
        hit = [g for p, g in groups.items() if cond(p)]
        if len(hit) != n_expect:
            raise ValueError(f"新形式: {key} の列グループが {len(hit)}個（想定 {n_expect}）")
        sell = buy = 0.0
        for g in hit:
            s_, b_, d_ = (_num(target[g[k]]) if k in g else None for k in ("売", "買", "差引"))
            if s_ is None or b_ is None or d_ is None:
                raise ValueError(f"新形式: {key} の売/買/差引が数値でない")
            # 自己検証: 差引 = 買 − 売 になっていなければ列を取り違えている
            if abs((b_ - s_) - d_) > 1:
                raise ValueError(f"新形式: {key} 差引が 買−売 と合わない（{b_}-{s_}≠{d_}）")
            sell += s_
            buy += b_
        flows[key] = {"sell": sell, "buy": buy}
        net[key] = round((buy - sell) / 100000)        # 千円 → 億円

    # 単位の確認: プライムの海外投資家の売りは週に数兆円＝10^9〜10^10千円台。桁が違えば単位が変わった
    if not (1e8 < flows["foreigners"]["sell"] < 1e12):
        raise ValueError(f"新形式: 金額の桁が想定外（海外投資家の売り {flows['foreigners']['sell']}）")

    return {"code": code, "market": market, "flows": flows, "net": net,
            # investor_hist.market_of はシート名で判定する。新形式はシートが1枚だけなので、
            # 市場名の列（B列）で東証プライムの行を選んだうえで、旧形式と同じ "TSE Prime" を入れる
            "sheet": "TSE Prime" if "TSE Prime" in market else market}


def new_week_meta(url_or_name, code):
    """新形式のファイル名（開始日・終了日）と A列の年月週コード（YYYYMMW）から
    開始日・終了日・旧形式と同じ書式の週ラベルを作る"""
    m = NEW_FILE_RE.search(url_or_name or "")
    if not m:
        return None, None, None
    d1 = datetime.strptime(m.group(1), "%Y%m%d").date()
    d2 = datetime.strptime(m.group(2), "%Y%m%d").date()
    label = None
    cm = re.fullmatch(r"(\d{4})(\d{2})(\d)", code or "")
    if cm:
        y, mo, w = int(cm.group(1)), int(cm.group(2)), int(cm.group(3))
        label = (f"{y}年{mo}月第{w}週 {y}/{mo} week{w} "
                 f"( {d1.month}/{d1.day} - {d2.month}/{d2.day} )")
    return d1.isoformat(), d2.isoformat(), label


# ===== 旧形式（〜2026年9月第2週 stock_val_1_YYMMWW.xls） =====
# 週ラベル行のセル「08/31～09/04」。前週・今週の2つが横に並ぶ（2016〜2026の557本すべて 10行目・D列とH列）
OLD_BLOCK_RE = re.compile(r"^\s*(\d{1,2})/(\d{1,2})\s*[～~〜]\s*(\d{1,2})/(\d{1,2})\s*$")


def old_label(df):
    """旧形式の週ラベル（例: 2026年9月第2週 2026/9 week2 ( 9/7 - 9/11 )）"""
    for i in range(min(8, len(df))):
        for v in df.iloc[i].tolist():
            s = str(v)
            if "週" in s and "年" in s:
                return re.sub(r"\s+", " ", s).strip()
    return None


def old_week_block(df):
    """週ラベル行を探し、今週ブロック（一番右）の 金額列・差引列・月日 を返す。

    見出し「金額 Value」は2列の結合セルで、数字は右隣の列に入っている（D列見出し→E列に値）。
    なので金額は「金額」見出しの列から「比率」見出しの手前までを範囲（amount_cols）で持つ。
    戻り値: {"row","amount_cols":range,"balance_col","md":(m1,d1,m2,d2),"blocks":ブロック数}
    """
    for i in range(min(20, len(df))):
        hits = []
        for c, v in enumerate(df.iloc[i].tolist()):
            m = OLD_BLOCK_RE.match(str(v))
            if m:
                hits.append((c, tuple(int(x) for x in m.groups())))
        if not hits:
            continue
        c0, md = max(hits)                    # 一番右 = 今週（左は前週）
        # 直下の段「金額 Value | 比率 Ratio | 差引き Balance」から、今週ブロック内の列を決める
        sub = [str(v) for v in df.iloc[i + 1].tolist()] if i + 1 < len(df) else []
        amount = next((c for c in range(c0, len(sub)) if sub[c].startswith("金額")), None)
        ratio = next((c for c in range(c0, len(sub)) if sub[c].startswith("比率")), None)
        balance = next((c for c in range(c0, len(sub)) if sub[c].startswith("差引")), None)
        if amount is None or ratio is None or not amount < ratio:
            raise ValueError("旧形式: 今週ブロックの「金額」「比率」列が見つからない")
        return {"row": i, "amount_cols": range(amount, ratio), "balance_col": balance,
                "md": md, "blocks": len(hits)}
    raise ValueError("旧形式: 週ラベル行（MM/DD～MM/DD）が見つからない")


def parse_old_df(df, code=None):
    """旧形式の1シート（DataFrame, header=None）から **今週** の6主体を取る。

    code（ファイル名の YYMMWW）があれば、週ラベルの「YYYY年M月第W週」とも突き合わせる。
    戻り値: {"label","start","end","flows":{key:{"sell","buy"}}(千円),"net":{key:億円}}
    """
    label = old_label(df)
    lm = re.search(r"(\d{4})年(\d{1,2})月第(\d)週", label or "")
    if not lm:
        raise ValueError(f"旧形式: 週ラベルが読めない（{label}）")
    y, mo, wk = (int(x) for x in lm.groups())
    if code and (2000 + int(code[:2]), int(code[2:4]), int(code[4:6])) != (y, mo, wk):
        raise ValueError(f"旧形式: ファイル名 {code} と週ラベル {y}年{mo}月第{wk}週 が合わない")
    start, end = week_dates(label, f"{y % 100:02d}")

    blk = old_week_block(df)
    m1, d1, m2, d2 = blk["md"]
    # ★ずれ再発防止★ 選んだブロックの日付 = ファイルの週（ラベル末尾の ( 9/7 - 9/11 )）
    if not start or (int(start[5:7]), int(start[8:10]), int(end[5:7]), int(end[8:10])) != (m1, d1, m2, d2):
        raise ValueError(f"旧形式: 今週ブロック {m1:02d}/{d1:02d}～{m2:02d}/{d2:02d} が"
                         f"ファイルの週 {start}〜{end} と合わない")

    ac, bc = blk["amount_cols"], blk["balance_col"]

    def amount(i):
        """今週ブロックの金額（金額見出し〜比率見出しの手前で最初の数値）"""
        return next((x for x in (_num(df.iloc[i, c]) for c in ac) if x is not None), None)

    flows, net = {}, {}
    for i in range(blk["row"] + 2, len(df) - 1):
        lab = str(df.iloc[i, 0]).replace("　", "").replace(" ", "")
        for prefix, key, _ in CATEGORIES:
            if key in flows or not lab.startswith(prefix):
                continue
            sell, buy = amount(i), amount(i + 1)
            if sell is None or buy is None:
                continue
            # 自己検証: 差引き（売り/買いどちらかの行に入る）= 買 − 売
            if bc is not None:
                bal = [b for b in (_num(df.iloc[i, bc]), _num(df.iloc[i + 1, bc])) if b is not None]
                if bal and abs((buy - sell) - bal[0]) > 1:
                    raise ValueError(f"旧形式: {key} 差引が 買−売 と合わない（{buy}-{sell}≠{bal[0]}）")
            flows[key] = {"sell": sell, "buy": buy}
            net[key] = round((buy - sell) / 100000)   # 千円 → 億円
    return {"label": label, "start": start, "end": end, "flows": flows, "net": net}


def read_old_sheet(raw, sheet=None):
    """旧形式 xls から 東証1部/プライム のシートを DataFrame で返す → (df, シート名)"""
    import pandas as pd
    xl = pd.ExcelFile(io.BytesIO(raw))
    sh = sheet or sheet_for(xl)
    return xl.parse(sheet_name=sh, header=None), sh


def parse(raw, code, url=""):
    if is_new_format(raw):
        w = parse_new(raw)
        start, end, label = new_week_meta(url, w["code"])
        return {"code": code, "label": label, "start": start, "end": end,
                "sheet": w["sheet"], "net": w["net"]}
    df, sh = read_old_sheet(raw)
    w = parse_old_df(df, code)
    return {"code": code, "label": w["label"], "start": w["start"], "end": w["end"],
            "sheet": sh, "net": w["net"]}


def main():
    CACHE.mkdir(exist_ok=True)
    print("JPXの週次ファイルを集める")
    urls = links()
    print(f"合計 {len(urls)}週ぶんのファイル")

    weeks, bad = [], []
    for i, code in enumerate(sorted(urls), 1):
        f = CACHE / (f"{code}.xlsx" if urls[code].endswith(".xlsx") else f"{code}.xls")
        if not f.exists():
            try:
                f.write_bytes(http(urls[code]))
            except Exception as e:
                bad.append((code, str(e)[:60]))
                continue
            time.sleep(0.4)          # JPXへの礼儀
        try:
            w = parse(f.read_bytes(), code, urls[code])
        except Exception as e:
            bad.append((code, str(e)[:60]))
            continue
        if len(w["net"]) < 4 or not w["start"]:
            bad.append((code, f"抽出不足 {len(w['net'])}主体 / start={w['start']}"))
            continue
        w["src"] = urls[code]
        weeks.append(w)
        if i % 60 == 0:
            print(f"  {i}/{len(urls)} 済み（採用{len(weeks)}・除外{len(bad)}）", flush=True)

    weeks.sort(key=lambda w: w["start"])
    out = {
        "updated": datetime.now(JST).isoformat(timespec="seconds"),
        "source": "JPX 投資部門別売買状況（週間・金額・委託と自己の合計）",
        "source_index": BASE + INDEX,
        "unit": "億円",
        "market_note": ("2022年4月4日の市場再編で対象が「東証1部」から「プライム」に変わりました。"
                        "構成銘柄が違うため、同じ系列として扱う場合は注意が必要です。"),
        "market_change": REORG,
        "categories": [{"key": k, "label": d} for _, k, d in CATEGORIES],
        "coverage": {"weeks": len(weeks),
                     "from": weeks[0]["start"] if weeks else None,
                     "to": weeks[-1]["end"] if weeks else None,
                     "skipped": len(bad),
                     "note": "JPXのアーカイブは2016年までしか遡れません。それ以前は取得していません"},
        # 区分は日付ではなくシート名で決める（次に再編があってもここは壊れない）
        "weeks": [{"s": w["start"], "e": w["end"], "m": market_of(w["sheet"]),
                   "net": w["net"]} for w in weeks],
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    print(f"\n採用 {len(weeks)}週 / 除外 {len(bad)}週")
    if bad:
        for c, e in bad[:10]:
            print(f"  除外 {c}: {e}")
    if weeks:
        print(f"期間 {weeks[0]['start']} 〜 {weeks[-1]['end']}")
        print(f"区分 東証1部 {sum(1 for w in out['weeks'] if w['m']=='tse1')}週 / "
              f"プライム {sum(1 for w in out['weeks'] if w['m']=='prime')}週")
    print(f"→ {OUT.name} ({OUT.stat().st_size/1024:.1f}KB)")


if __name__ == "__main__":
    main()
